import io
import os
import shutil
import subprocess
import time

import pytest

from conftest import H, SAMPLE, upload
from app import ai, comm, extract

FFMPEG = shutil.which("ffmpeg")


# ------------------------------------------------------------------ auth & health
def test_health(client):
    r = client.get("/api/health")
    assert r.status_code == 200
    assert {"ai", "ffmpeg", "data_dir_writable"} <= set(r.json())
    assert r.json()["data_dir_writable"] is True


def test_second_admin_self_register_blocked(client, w):
    r = client.post("/api/auth/register", json={"name": "X", "email": "x@t.com", "password": "pw123456", "role": "admin"})
    assert r.status_code == 403


def test_duplicate_email_and_bad_login(client, w):
    assert client.post("/api/auth/register", json={"name": "D", "email": "s1@t.com", "password": "pw123456", "role": "student"}).status_code == 400
    assert client.post("/api/auth/login", json={"email": "s1@t.com", "password": "wrong"}).status_code == 401
    assert client.post("/api/auth/login", json={"email": "nobody@t.com", "password": "x"}).status_code == 401


def test_me_requires_token(client, w):
    assert client.get("/api/me").status_code == 401
    assert client.get("/api/me", headers={"Authorization": "Bearer junk"}).status_code == 401
    me = client.get("/api/me", headers=w["h"]["s1"]).json()
    assert me["role"] == "student" and me["ai_enabled"] is False  # no key in tests


# ------------------------------------------------------------------ student restrictions
def test_student_cannot_create_content(client, w, shared_mat):
    s = w["h"]["s1"]
    assert upload(client, w, who="s1", text=SAMPLE).status_code == 403
    assert client.post("/api/classrooms", headers=s, json={"name": "x"}).status_code == 403
    assert client.post(f"/api/materials/{shared_mat}/quizzes", headers=s, json={"count": 5}).status_code == 403
    assert client.post(f"/api/materials/{shared_mat}/video", headers=s).status_code == 403
    assert client.delete(f"/api/materials/{shared_mat}", headers=s).status_code == 403
    assert client.post(f"/api/materials/{shared_mat}/share", headers=s, json={"classroom_id": None}).status_code == 403


def test_parent_sees_no_materials(client, w, shared_mat):
    assert client.get("/api/materials", headers=w["h"]["p1"]).json() == []


# ------------------------------------------------------------------ upload / sharing / visibility
def test_upload_validation(client, w):
    assert client.post("/api/materials", headers=w["h"]["t1"], data={"title": "", "text": ""}).status_code == 400
    assert upload(client, w, text=SAMPLE, classroom_id=w["cid2"]).status_code == 403  # another teacher's class
    assert upload(client, w, who="t1", filename="empty.txt", data=b"").status_code == 422  # nothing readable


def test_private_material_hidden_then_shared(client, w):
    r = upload(client, w, title="Private lesson", text=SAMPLE)
    assert r.status_code == 200
    m = r.json()
    assert m["summary"] and m["class_name"] == "" and m["title"] == "Private lesson"
    assert all(x["id"] != m["id"] for x in client.get("/api/materials", headers=w["h"]["s1"]).json())
    assert client.get(f"/api/materials/{m['id']}", headers=w["h"]["s1"]).status_code == 404
    # quiz on a private material is invisible too
    q = client.post(f"/api/materials/{m['id']}/quizzes", headers=w["h"]["t1"], json={"count": 4, "types": ["mcq", "tf"]}).json()
    assert client.get(f"/api/quizzes/{q['id']}", headers=w["h"]["s1"]).status_code == 404
    # other teacher can't share it
    assert client.post(f"/api/materials/{m['id']}/share", headers=w["h"]["t2"], json={"classroom_id": w["cid2"]}).status_code in (403, 404)
    # share with class -> material AND quiz appear for enrolled students only
    r = client.post(f"/api/materials/{m['id']}/share", headers=w["h"]["t1"], json={"classroom_id": w["cid"]})
    assert r.status_code == 200 and r.json()["class_name"] == "Biology 9"
    assert any(x["id"] == m["id"] for x in client.get("/api/materials", headers=w["h"]["s1"]).json())
    assert any(x["id"] == m["id"] for x in client.get("/api/materials", headers=w["h"]["s2"]).json())
    assert client.get("/api/materials", headers=w["h"]["s3"]).json() == [] or all(
        x["id"] != m["id"] for x in client.get("/api/materials", headers=w["h"]["s3"]).json())
    assert client.get(f"/api/quizzes/{q['id']}", headers=w["h"]["s1"]).status_code == 200
    # make private again
    client.post(f"/api/materials/{m['id']}/share", headers=w["h"]["t1"], json={"classroom_id": None})
    assert client.get(f"/api/materials/{m['id']}", headers=w["h"]["s1"]).status_code == 404


def test_student_view_hides_raw_text(client, w, shared_mat):
    d = client.get(f"/api/materials/{shared_mat}", headers=w["h"]["s1"]).json()
    assert "text" not in d and d["posted_by"] == "Teacher One" and d["class_name"] == "Biology 9"
    assert "text" in client.get(f"/api/materials/{shared_mat}", headers=w["h"]["t1"]).json()
    assert client.get(f"/api/materials/{shared_mat}", headers=w["h"]["s3"]).status_code == 404  # not enrolled


def test_upload_docx_and_pptx(client, w):
    import docx
    from pptx import Presentation
    d = docx.Document(); d.add_paragraph(SAMPLE); b = io.BytesIO(); d.save(b)
    r = upload(client, w, filename="lesson.docx", data=b.getvalue())
    assert r.status_code == 200 and r.json()["kind"] == "document" and "chlorophyll" in r.json()["text"].lower()
    p = Presentation(); s = p.slides.add_slide(p.slide_layouts[1]); s.shapes.title.text = "Photosynthesis"
    s.placeholders[1].text = SAMPLE[:200]; b = io.BytesIO(); p.save(b)
    r = upload(client, w, filename="slides.pptx", data=b.getvalue())
    assert r.status_code == 200 and "Photosynthesis" in r.json()["text"]


def test_image_without_working_ai_is_rejected_not_saved(client, w):
    from PIL import Image
    b = io.BytesIO(); Image.new("RGB", (60, 60), "white").save(b, "PNG")
    before = len(client.get("/api/materials", headers=w["h"]["t1"]).json())
    r = upload(client, w, filename="diagram.png", data=b.getvalue(), ctype="image/png")
    assert r.status_code == 422 and "AI" in r.json()["detail"]
    assert len(client.get("/api/materials", headers=w["h"]["t1"]).json()) == before  # no junk material created


def test_bad_api_key_is_reported_and_never_saved_as_content(client, w, monkeypatch):
    """The reported bug: invalid key (401) used to become the material's title and summary."""
    class Boom:
        class messages:
            @staticmethod
            def create(**k):
                raise Exception("Error code: 401 - {'type': 'error', 'error': {'type': 'authentication_error', 'message': 'invalid x-api-key'}}")
    monkeypatch.setenv("ANTHROPIC_API_KEY", "sk-ant-bogus-123")
    monkeypatch.setattr(ai, "client", lambda: Boom())
    ai._st.update(t=0, ok=False, error="")
    try:
        me = client.get("/api/me", headers=w["h"]["t1"]).json()
        assert me["ai_enabled"] is False and "rejected" in me["ai_error"]
        assert client.get("/api/health").json()["ai"]["working"] is False
        from PIL import Image
        b = io.BytesIO(); Image.new("RGB", (60, 60), "white").save(b, "PNG")
        r = upload(client, w, filename="d.png", data=b.getvalue(), ctype="image/png")
        assert r.status_code == 422 and "rejected" in r.json()["detail"]
        r = upload(client, w, title="Still works", text=SAMPLE)  # text lessons still work with basic summary
        assert r.status_code == 200 and "Could not" not in r.json()["title"] and "401" not in r.json()["summary"]
    finally:
        ai._st.update(t=0, ok=False, error="")


@pytest.mark.skipif(not FFMPEG, reason="ffmpeg not installed")
def test_video_upload_is_shared_even_without_transcript(client, w, tmp_path):
    mp4 = tmp_path / "lecture.mp4"
    subprocess.run(["ffmpeg", "-y", "-loglevel", "error", "-f", "lavfi", "-i", "testsrc=duration=1:size=160x120:rate=10",
                    "-f", "lavfi", "-i", "sine=duration=1", "-shortest", str(mp4)], check=True)
    r = upload(client, w, title="", filename="lecture.mp4", data=mp4.read_bytes(), ctype="video/mp4", classroom_id=w["cid"])
    assert r.status_code == 200, r.text
    m = r.json()
    assert m["kind"] == "video" and m["has_file"] and m["title"] == "lecture.mp4" and m["warning"]
    # students can download/stream the teacher's original video
    f = client.get(f"/api/materials/{m['id']}/file", params={"token": w["tok"]["s1"]})
    assert f.status_code == 200 and len(f.content) > 1000
    assert client.get(f"/api/materials/{m['id']}/file", params={"token": w["tok"]["s3"]}).status_code == 404
    assert client.get(f"/api/materials/{m['id']}/file").status_code == 401
    # no text -> can't build an explainer video from it
    assert client.post(f"/api/materials/{m['id']}/video", headers=w["h"]["t1"]).status_code == 422


# ------------------------------------------------------------------ summaries, tutor
def test_summarize_and_tutor_offline(client, w, shared_mat):
    s = w["h"]["s1"]
    assert client.post(f"/api/materials/{shared_mat}/summarize", headers=s, json={"lang": "en"}).json()["summary"]
    assert client.post(f"/api/materials/{shared_mat}/summarize", headers=s, json={"lang": "te"}).status_code == 503
    r = client.post("/api/tutor/chat", headers=s, json={"message": "what absorbs sunlight?", "lang": "en", "material_id": shared_mat})
    assert r.status_code == 200 and "class material" in r.json()["reply"].lower()
    r = client.post("/api/tutor/chat", headers=s, json={"message": "chlorophyll ante enti?", "lang": "te_en", "material_id": shared_mat})
    assert r.status_code == 200 and "ANTHROPIC_API_KEY" in r.json()["reply"]
    assert client.post("/api/tutor/chat", headers=s, json={"message": "  "}).status_code == 400
    assert client.post("/api/tutor/chat", headers=w["h"]["s3"], json={"message": "hi", "material_id": shared_mat}).status_code == 404


# ------------------------------------------------------------------ quizzes, scoring, progress, feedback
def test_quiz_flow(client, w, quiz):
    assert len(quiz["questions"]) >= 3 and all("answer" in q for q in quiz["questions"])  # teacher sees answers
    lst = client.get("/api/quizzes", headers=w["h"]["s1"]).json()
    assert any(q["id"] == quiz["id"] for q in lst)
    sv = client.get(f"/api/quizzes/{quiz['id']}", headers=w["h"]["s1"]).json()
    assert all("answer" not in q and "explanation" not in q for q in sv["questions"])  # students don't
    assert client.get(f"/api/quizzes/{quiz['id']}", headers=w["h"]["s3"]).status_code == 404

    blank = client.post(f"/api/quizzes/{quiz['id']}/submit", headers=w["h"]["s1"], json={"answers": {}, "seconds": 10}).json()
    assert blank["pct"] == 0
    answers = {str(q["id"]): q["answer"] for q in quiz["questions"]}
    full = client.post(f"/api/quizzes/{quiz['id']}/submit", headers=w["h"]["s1"], json={"answers": answers, "seconds": 42}).json()
    assert full["pct"] == 100 and full["score"] == full["total"]
    assert client.post(f"/api/quizzes/{quiz['id']}/submit", headers=w["h"]["t1"], json={"answers": {}}).status_code == 403  # teachers can't submit
    assert client.post(f"/api/quizzes/{quiz['id']}/submit", headers=w["h"]["s3"], json={"answers": {}}).status_code == 404


def test_progress_visibility_and_feedback(client, w):
    sid = w["id"]["s1"]
    p = client.get(f"/api/students/{sid}/progress", headers=w["h"]["s1"]).json()
    assert p["stats"]["attempts"] >= 2 and 0 <= p["stats"]["average"] <= 100 and p["stats"]["daily"]
    assert client.get(f"/api/students/{sid}/progress", headers=w["h"]["t1"]).status_code == 200      # their teacher
    assert client.get(f"/api/students/{sid}/progress", headers=w["h"]["p1"]).status_code == 200      # linked parent
    assert client.get(f"/api/students/{sid}/progress", headers=w["h"]["admin"]).status_code == 200
    assert client.get(f"/api/students/{sid}/progress", headers=w["h"]["t2"]).status_code == 403      # other teacher
    assert client.get(f"/api/students/{sid}/progress", headers=w["h"]["s2"]).status_code == 403      # other student
    assert client.get(f"/api/students/{w['id']['s2']}/progress", headers=w["h"]["p1"]).status_code == 403  # unlinked child
    assert client.post(f"/api/students/{sid}/feedback/ai", headers=w["h"]["s1"]).status_code == 200
    assert client.post(f"/api/students/{sid}/feedback", headers=w["h"]["t1"], json={"text": "Good work!"}).status_code == 200
    assert client.post(f"/api/students/{sid}/feedback", headers=w["h"]["s1"], json={"text": "x"}).status_code == 403
    fb = client.get(f"/api/students/{sid}/progress", headers=w["h"]["p1"]).json()["feedback"]
    assert {f["by"] for f in fb} >= {"AI tutor", "Teacher One"}
    assert client.post(f"/api/students/{w['id']['s2']}/feedback/ai", headers=w["h"]["s2"]).status_code == 422  # no attempts yet


# ------------------------------------------------------------------ explainer video
@pytest.mark.skipif(not FFMPEG, reason="ffmpeg not installed")
def test_explainer_video_generation(client, w, shared_mat):
    t1 = w["h"]["t1"]
    r = client.post(f"/api/materials/{shared_mat}/video", headers=t1)
    assert r.status_code == 200 and r.json()["video_status"] in ("processing", "ready")
    for _ in range(120):
        m = client.get(f"/api/materials/{shared_mat}", headers=t1).json()
        if m["video_status"] in ("ready", "failed"):
            break
        time.sleep(1)
    assert m["video_status"] == "ready", m["video_error"]
    v = client.get(f"/api/materials/{shared_mat}/video", params={"token": w["tok"]["s1"]})
    assert v.status_code == 200 and v.headers["content-type"] == "video/mp4" and len(v.content) > 5000
    assert client.get(f"/api/materials/{shared_mat}/video", params={"token": w["tok"]["s3"]}).status_code == 404
    assert client.get(f"/api/materials/{shared_mat}/video").status_code == 401


# ------------------------------------------------------------------ communicator
def test_daily_tasks_differ_by_level_and_are_stable(client, w):
    a = client.get("/api/comm/today", headers=w["h"]["s1"]).json()      # high school
    b = client.get("/api/comm/today", headers=w["h"]["s2"]).json()      # primary
    assert a["level"] == "high school" and b["level"] == "primary"
    assert [t["skill"] for t in a["tasks"]] == ["speak", "write", "vocab"]
    assert a["tasks"][0]["target"]["seconds"] == 60 and b["tasks"][0]["target"]["seconds"] == 30
    assert a["tasks"][1]["target"]["min_words"] > b["tasks"][1]["target"]["min_words"]
    assert a["tasks"][0]["title"] != b["tasks"][0]["title"]
    again = client.get("/api/comm/today", headers=w["h"]["s1"]).json()
    assert [t["id"] for t in again["tasks"]] == [t["id"] for t in a["tasks"]]  # not regenerated


def test_task_submission_rules_and_feedback(client, w):
    tasks = client.get("/api/comm/today", headers=w["h"]["s1"]).json()["tasks"]
    tid = tasks[1]["id"]
    assert client.post(f"/api/comm/tasks/{tid}/submit", headers=w["h"]["s1"], json={"response": "hi"}).status_code == 400
    assert client.post(f"/api/comm/tasks/{tid}/submit", headers=w["h"]["s2"], json={"response": SAMPLE}).status_code == 404  # not theirs
    r = client.post(f"/api/comm/tasks/{tid}/submit", headers=w["h"]["s1"], json={"response": SAMPLE, "seconds": 0})
    assert r.status_code == 200, r.text
    t = r.json()["task"]
    assert t["status"] == "done" and t["xp"] > 0 and {"clarity", "grammar", "vocabulary", "structure"} <= set(t["scores"])
    assert all(1 <= v <= 5 for v in t["scores"].values()) and t["feedback"]["strengths"] and t["feedback"]["improvements"]
    s = r.json()["summary"]
    assert s["streak"] == 1 and s["today_done"] == 1 and s["xp"] == t["xp"]
    assert client.post(f"/api/comm/tasks/{tid}/submit", headers=w["h"]["s1"], json={"response": SAMPLE}).status_code == 409
    # speaking: pace and fillers are measured
    sp = tasks[0]["id"]
    r = client.post(f"/api/comm/tasks/{sp}/submit", headers=w["h"]["s1"],
                    json={"response": "Um I think, like, exams are stressful but planning helps. You know, short breaks help too.", "seconds": 10})
    m = r.json()["task"]["feedback"]["metrics"]
    assert m["fillers"] >= 3 and m["wpm"] > 0 and "fluency" in r.json()["task"]["scores"]


def test_teacher_assigned_task_and_comm_in_progress_view(client, w):
    before = len(client.get("/api/comm/today", headers=w["h"]["s1"]).json()["tasks"])
    body = {"classroom_id": w["cid"], "title": "Introduce yourself", "prompt": "Introduce yourself in 5 sentences.", "skill": "speak", "min_words": 30}
    assert client.post("/api/comm/assign", headers=w["h"]["t2"], json=body).status_code == 404       # not their class
    assert client.post("/api/comm/assign", headers=w["h"]["s1"], json=body).status_code == 403
    assert client.post("/api/comm/assign", headers=w["h"]["t1"], json={**body, "skill": "dance"}).status_code == 400
    assert client.post("/api/comm/assign", headers=w["h"]["t1"], json=body).json()["assigned_to"] == 2
    after = client.get("/api/comm/today", headers=w["h"]["s1"]).json()["tasks"]
    assert len(after) == before + 1 and any(t["source"] == "teacher" for t in after)
    assert len(client.get("/api/comm/today", headers=w["h"]["s2"]).json()["tasks"]) == 4  # 3 daily + 1 teacher
    prog = client.get(f"/api/students/{w['id']['s1']}/progress", headers=w["h"]["t1"]).json()
    assert prog["communication"]["completed"] >= 2 and prog["communication"]["criteria"]


def test_roleplay_scenarios(client, w):
    sc = client.get("/api/comm/scenarios", headers=w["h"]["s1"]).json()
    assert {s["id"] for s in sc} == {"council", "admission", "debate"}
    assert {s["id"] for s in client.get("/api/comm/scenarios", headers=w["h"]["s2"]).json()} == {"friend", "canteen", "weekend"}
    r = client.post("/api/comm/chat", headers=w["h"]["s1"], json={"scenario_id": "council", "message": "", "history": []})
    assert r.status_code == 200 and r.json()["reply"]
    r = client.post("/api/comm/chat", headers=w["h"]["s1"], json={"scenario_id": "council", "message": "I want to improve the library.",
                                                                  "history": [{"role": "assistant", "content": "Hello"}]})
    assert "Coach:" in r.json()["reply"]
    assert client.post("/api/comm/chat", headers=w["h"]["s1"], json={"scenario_id": "nope", "message": "x"}).status_code == 404
    assert client.get("/api/comm/today", headers=w["h"]["p1"]).status_code == 403  # parents have no tasks


# ------------------------------------------------------------------ admin & classes
def test_admin_and_classes(client, w):
    assert client.get("/api/admin/users", headers=w["h"]["t1"]).status_code == 403
    assert client.get("/api/admin/users", headers=w["h"]["s1"]).status_code == 403
    ov = client.get("/api/admin/overview", headers=w["h"]["admin"]).json()
    assert ov["users"]["student"] == 3 and ov["users"]["teacher"] == 2 and ov["materials"] >= 1
    assert client.delete(f"/api/admin/users/{w['id']['admin']}", headers=w["h"]["admin"]).status_code == 400
    r = client.post("/api/admin/users", headers=w["h"]["admin"], json={"name": "Tmp", "email": "tmp@t.com", "password": "pw123456", "role": "student"})
    assert r.status_code == 200
    assert client.post("/api/admin/users", headers=w["h"]["admin"], json={"name": "Tmp", "email": "tmp@t.com", "password": "pw123456", "role": "student"}).status_code == 400
    assert client.post("/api/admin/users", headers=w["h"]["admin"], json={"name": "T", "email": "bad@t.com", "password": "pw123456", "role": "wizard"}).status_code == 400
    assert client.delete(f"/api/admin/users/{r.json()['id']}", headers=w["h"]["admin"]).status_code == 200
    assert client.post("/api/auth/login", json={"email": "tmp@t.com", "password": "pw123456"}).status_code == 401
    cls = client.get("/api/classrooms", headers=w["h"]["s1"]).json()
    assert [c["name"] for c in cls] == ["Biology 9"]
    assert client.get("/api/classrooms", headers=w["h"]["s3"]).json() == []
    assert client.post(f"/api/classrooms/{w['cid']}/students", headers=w["h"]["t1"], json={"student_email": "ghost@t.com"}).status_code == 404
    assert client.post(f"/api/classrooms/{w['cid']}/students", headers=w["h"]["t1"], json={"student_email": "t2@t.com"}).status_code == 404  # not a student
    rows = client.get(f"/api/classrooms/{w['cid']}/progress", headers=w["h"]["t1"]).json()
    assert {r["student"]["email"] for r in rows} == {"s1@t.com", "s2@t.com"}
    assert client.get(f"/api/classrooms/{w['cid']}/progress", headers=w["h"]["t2"]).status_code == 404


def test_delete_material_cascades(client, w):
    m = upload(client, w, text=SAMPLE, classroom_id=w["cid"]).json()
    q = client.post(f"/api/materials/{m['id']}/quizzes", headers=w["h"]["t1"], json={"count": 4, "types": ["mcq", "tf"]}).json()
    assert client.delete(f"/api/materials/{m['id']}", headers=w["h"]["t2"]).status_code in (403, 404)
    assert client.delete(f"/api/materials/{m['id']}", headers=w["h"]["t1"]).status_code == 200
    assert client.get(f"/api/materials/{m['id']}", headers=w["h"]["t1"]).status_code == 404
    assert client.get(f"/api/quizzes/{q['id']}", headers=w["h"]["t1"]).status_code == 404


# ------------------------------------------------------------------ role x endpoint permission matrix
def test_unauthenticated_requests_are_rejected(client, w, shared_mat):
    for method, url in [("GET", "/api/me"), ("GET", "/api/materials"), ("GET", f"/api/materials/{shared_mat}"), ("GET", "/api/quizzes"),
                        ("GET", "/api/admin/users"), ("GET", "/api/classrooms"), ("GET", "/api/comm/today"),
                        ("POST", "/api/tutor/chat"), ("GET", f"/api/students/{w['id']['s1']}/progress")]:
        assert client.request(method, url).status_code == 401, (method, url)


def test_role_permission_matrix(client, w, shared_mat):
    cid, s1 = w["cid"], w["id"]["s1"]
    #  method, url, json body, {role: expected status}
    matrix = [
        ("GET", "/api/admin/users", None, {"admin": 200, "t1": 403, "s1": 403, "p1": 403}),
        ("GET", "/api/admin/overview", None, {"admin": 200, "t1": 403, "s1": 403, "p1": 403}),
        ("POST", "/api/admin/parent-links", {"parent_email": "p1@t.com", "student_email": "s1@t.com"}, {"admin": 200, "t1": 200, "s1": 403, "p1": 403}),
        ("POST", "/api/classrooms", {"name": "Matrix class"}, {"admin": 200, "t1": 200, "s1": 403, "p1": 403}),
        ("GET", "/api/materials", None, {"admin": 200, "t1": 200, "s1": 200, "p1": 200}),
        ("GET", "/api/quizzes", None, {"admin": 200, "t1": 200, "s1": 200, "p1": 200}),
        ("GET", "/api/parent/children", None, {"p1": 200, "admin": 403, "t1": 403, "s1": 403}),
        ("GET", "/api/comm/today", None, {"admin": 200, "t1": 200, "s1": 200, "p1": 403}),
        ("POST", "/api/tutor/chat", {"message": "hello"}, {"admin": 200, "t1": 200, "s1": 200, "p1": 403}),
        ("GET", f"/api/classrooms/{cid}/progress", None, {"admin": 200, "t1": 200, "t2": 404, "s1": 403, "p1": 403}),
        ("POST", f"/api/classrooms/{cid}/students", {"student_email": "s1@t.com"}, {"admin": 200, "t1": 200, "t2": 404, "s1": 403, "p1": 403}),
        ("POST", f"/api/materials/{shared_mat}/summarize", {"lang": "en"}, {"admin": 200, "t1": 200, "s1": 200, "s3": 404, "p1": 404}),
        ("POST", f"/api/materials/{shared_mat}/quizzes", {"count": 3, "types": ["mcq", "tf"]}, {"s1": 403, "p1": 403, "t2": 404}),
        ("POST", f"/api/materials/{shared_mat}/video", None, {"s1": 403, "p1": 403, "t2": 404}),
        ("POST", f"/api/materials/{shared_mat}/share", {"classroom_id": cid}, {"s1": 403, "p1": 403, "t2": 404, "t1": 200}),
        ("GET", f"/api/students/{s1}/progress", None, {"s1": 200, "t1": 200, "p1": 200, "admin": 200, "t2": 403, "s2": 403, "s3": 403}),
    ]
    bad = []
    for method, url, body, expect in matrix:
        for role, code in expect.items():
            r = client.request(method, url, headers=w["h"][role], json=body)
            if r.status_code != code:
                bad.append(f"{role} {method} {url} -> {r.status_code}, expected {code}")
    assert not bad, "\n".join(bad)


def test_parent_children_and_assessment_kind(client, w):
    kids = client.get("/api/parent/children", headers=w["h"]["p1"]).json()
    assert [k["email"] for k in kids] == ["s1@t.com"]
    m = upload(client, w, title="Cells", text=SAMPLE, classroom_id=w["cid"]).json()
    q = client.post(f"/api/materials/{m['id']}/quizzes", headers=w["h"]["t1"], json={"count": 4, "kind": "assessment", "types": ["mcq", "tf"]}).json()
    assert q["kind"] == "assessment" and q["title"].startswith("Assessment:")
    row = [x for x in client.get("/api/quizzes", headers=w["h"]["s1"]).json() if x["id"] == q["id"]][0]
    assert row["best_pct"] is None  # not attempted yet


# ------------------------------------------------------------------ pure unit tests (no server)
def test_unit_ai_fallbacks():
    assert ai.status()["configured"] is False and ai.enabled() is False
    a = ai.analyze(SAMPLE, "high school")
    assert a["summary"] and a["key_points"] and a["topics"] and "_fallback" not in a
    qs = ai.make_questions(SAMPLE, "high school", 5, ["mcq", "tf"], a["topics"])
    assert len(qs) >= 3 and all(q["id"] and q["answer"] and q["type"] in ("mcq", "tf") for q in qs)
    assert all(q["answer"] in q["options"] for q in qs)
    s, _ = ai.grade_short("q", "chlorophyll absorbs sunlight", "it absorbs sunlight using chlorophyll")
    assert s > 0.5
    assert ai.grade_short("q", "x", "")[0] == 0
    assert "rejected" in ai.friendly(Exception("Error code: 401 - invalid x-api-key"))


def test_unit_extract_and_comm():
    assert extract.detect_kind("a.PNG") == "image" and extract.detect_kind("a.mp4") == "video" and extract.detect_kind("a.pdf") == "document"
    m = comm.metrics("Um, I like, you know, basically agree. It is good.", 6)
    assert m["fillers"] >= 4 and m["sentences"] == 2 and m["wpm"] > 0
    assert comm.band("unknown level") == "high school"
    for lvl, b in comm.BANK.items():
        assert len(b["topics"]) >= 3 and len(b["writes"]) >= 3 and len(b["vocab"]) >= 3 and lvl in comm.SCENARIOS
    assert comm.BANK["primary"]["secs"] < comm.BANK["high school"]["secs"] < comm.BANK["postgraduate"]["secs"]
