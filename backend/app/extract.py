"""Turn any uploaded file (doc, image, audio, video, text) into plain text for analysis."""
import os
import shutil
import subprocess
import tempfile
from . import ai

DOC_EXT = {".pdf", ".docx", ".pptx", ".txt", ".md", ".csv", ".html", ".htm", ".rtf", ".json"}
IMG_EXT = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp"}
AUDIO_EXT = {".mp3", ".wav", ".m4a", ".ogg", ".flac"}
VIDEO_EXT = {".mp4", ".mov", ".mkv", ".avi", ".webm"}


def detect_kind(filename: str) -> str:
    ext = os.path.splitext(filename.lower())[1]
    if ext in IMG_EXT:
        return "image"
    if ext in AUDIO_EXT:
        return "audio"
    if ext in VIDEO_EXT:
        return "video"
    return "document"


def extract_text(path: str, filename: str) -> tuple[str, str]:
    """Returns (kind, text). Never raises for unsupported content; returns a note instead."""
    ext = os.path.splitext(filename.lower())[1]
    kind = detect_kind(filename)
    try:
        if kind == "image":
            return kind, _image(path, ext)
        if kind == "audio":
            return kind, _transcribe(path)
        if kind == "video":
            return kind, _video(path)
        return kind, _document(path, ext)
    except Exception as e:  # noqa
        return kind, f"[Could not extract content: {e}]"


def _document(path, ext):
    if ext == ".pdf":
        from pypdf import PdfReader
        r = PdfReader(path)
        return "\n".join((p.extract_text() or "") for p in r.pages)
    if ext == ".docx":
        import docx
        d = docx.Document(path)
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts)
    if ext == ".pptx":
        from pptx import Presentation
        out = []
        for i, s in enumerate(Presentation(path).slides, 1):
            out.append(f"Slide {i}")
            for sh in s.shapes:
                if sh.has_text_frame:
                    out.append(sh.text_frame.text)
        return "\n".join(out)
    with open(path, "rb") as f:
        raw = f.read()
    txt = raw.decode("utf-8", errors="ignore")
    if ext in (".html", ".htm"):
        import re
        txt = re.sub(r"<(script|style).*?</\1>", " ", txt, flags=re.S)
        txt = re.sub(r"<[^>]+>", " ", txt)
    return txt


def _image(path, ext):
    # Prefer Claude vision; fall back to OCR if available.
    text = ai.describe_image(path, ext)
    if text:
        return text
    try:
        import pytesseract
        from PIL import Image
        return pytesseract.image_to_string(Image.open(path))
    except Exception:
        return "[Image uploaded, but no vision model or OCR is configured. Set ANTHROPIC_API_KEY to analyse images.]"


def _transcribe(path):
    try:
        from faster_whisper import WhisperModel
    except ImportError:
        return "[Audio uploaded, but transcription is not installed. Run: pip install faster-whisper]"
    model = WhisperModel(os.environ.get("WHISPER_MODEL", "base"), compute_type="int8")
    segs, _ = model.transcribe(path)
    return " ".join(s.text.strip() for s in segs)


def _video(path):
    if not shutil.which("ffmpeg"):
        return "[Video uploaded, but ffmpeg is not installed, so audio could not be extracted.]"
    with tempfile.TemporaryDirectory() as td:
        wav = os.path.join(td, "a.wav")
        subprocess.run(["ffmpeg", "-y", "-i", path, "-vn", "-ac", "1", "-ar", "16000", wav],
                       capture_output=True, check=True)
        return _transcribe(wav)
